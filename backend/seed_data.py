import os
import sys
from pathlib import Path
import openpyxl
from backend.database import get_db_connection, init_db

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

BASE_DIR = Path(__file__).resolve().parent.parent
SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events() -> int:
    """
    Seeds the events table from data/events_seed.xlsx.
    Raises FileNotFoundError if the companion seed file is missing.
    Skips seeding if the events table already has rows.
    Returns the number of seeded rows (or 0 if already seeded).
    """
    if not SEED_FILE_PATH.exists():
        raise FileNotFoundError(
            f"Seed data file not found at '{SEED_FILE_PATH}'. "
            "Please copy the provided 'events_seed.xlsx' file into the 'data/' folder before running this script."
        )

    init_db()
    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if events table already has data
    cursor.execute("SELECT COUNT(*) AS count FROM events;")
    existing_count = cursor.fetchone()["count"]
    if existing_count > 0:
        conn.close()
        return 0

    wb = openpyxl.load_workbook(SEED_FILE_PATH, data_only=True)
    if "Events" not in wb.sheetnames:
        conn.close()
        raise ValueError(
            f"The Excel file '{SEED_FILE_PATH}' does not contain an 'Events' sheet."
        )

    sheet = wb["Events"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        conn.close()
        raise ValueError("The 'Events' sheet is empty.")

    header = [str(col).strip().lower() if col is not None else "" for col in rows[0]]
    expected_header = [
        "id",
        "title",
        "category",
        "description",
        "event_date",
        "event_time",
        "venue",
        "capacity",
        "fee",
        "icon",
    ]

    # Insert rows
    inserted = 0
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        row_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = (
            row[:10]
        )
        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """,
            (
                int(row_id),
                str(title),
                str(category),
                str(description) if description is not None else None,
                str(event_date),
                str(event_time),
                str(venue),
                int(capacity),
                int(fee) if fee is not None else 0,
                str(icon) if icon is not None else None,
            ),
        )
        inserted += 1

    conn.commit()
    conn.close()
    return inserted


if __name__ == "__main__":
    count = seed_events()
    print(f"Seeding completed. Rows inserted: {count}")
    conn = get_db_connection()
    rows = conn.execute("SELECT * FROM events ORDER BY id ASC;").fetchall()
    print(f"Total events in database: {len(rows)}")
    for r in rows:
        print(f"[{r['id']}] {r['icon']} {r['title']} ({r['category']}) - {r['event_date']} {r['event_time']} @ {r['venue']} | Cap: {r['capacity']} | Fee: {r['fee']}")
    conn.close()
